"""Read existing reports only. Never writes to or triggers either source repository."""
import os, io, json, re, shutil, zipfile
from pathlib import Path
from datetime import datetime, timezone
import requests
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent
OUT=ROOT/'site'
DNL='Medtech-platform/DNLAutomation'
RADAR='Medtech-platform/News-Radar'
PUBLIC='https://medtech-platform.github.io/News-Radar/data/'

def get(url, token=None, redirect=True):
    headers={'Accept':'application/vnd.github+json','X-GitHub-Api-Version':'2022-11-28'}
    if token: headers['Authorization']='Bearer '+token
    r=requests.get(url,headers=headers,timeout=90,allow_redirects=redirect)
    if r.status_code not in (200,302):
        raise RuntimeError(f'Report read failed (HTTP {r.status_code}). Check source access and the INTELHUB_READ_TOKEN secret.')
    return r

def save_json(path,data):
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str),encoding='utf-8')

def convert_excel(data):
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    ws=wb['DNL Report'] if 'DNL Report' in wb.sheetnames else wb.active
    rows=ws.iter_rows(values_only=True)
    headers=[str(x or '').strip() for x in next(rows)]
    mapping={'Headline':'headline','Company Name':'company','Screened Entity':'screened_entity','Screening Category':'screening_category','Date':'date','News Type':'news_type','Source Link':'url','Source Type':'source_type','Hot vs Non-Hot':'hot','Date Collected':'date_collected'}
    if 'Headline' not in headers or 'Source Link' not in headers:raise ValueError('DNL report columns have changed. Expected Headline and Source Link.')
    result=[]
    for row in rows:
        d=dict(zip(headers,row))
        if d.get('Headline'):
            result.append({v: str(d.get(k) or '') for k,v in mapping.items()})
    wb.close()
    return result

def sync_radar():
    target=OUT/'data';target.mkdir(parents=True,exist_ok=True)
    files=get(f'https://api.github.com/repos/{RADAR}/contents/docs/data').json()
    if not isinstance(files,list):raise ValueError('News Radar directory could not be read.')
    names={f['name'] for f in files};runs=[]
    for name in sorted(names,reverse=True):
        match=re.fullmatch(r'rxbenefits_(\d{4}-\d{2}-\d{2})\.json',name)
        if not match:continue
        report=get(PUBLIC+name).json()
        if not isinstance(report.get('articles'),list):raise ValueError('Unexpected News Radar data format')
        save_json(target/name,report)
        day=match[1];excel=f'rxbenefits_intel_hub_report_{day}.xlsx'
        if excel in names:(target/excel).write_bytes(get(PUBLIC+excel).content)
        else:excel=None
        runs.append({'date':day,'file':name,'count':len(report['articles']),'excel':excel})
    if not runs:raise ValueError('No published News Radar reports were found.')
    save_json(target/'index.json',{'runs':runs})
    print(f'Read {len(runs)} existing News Radar reports')

def sync_dnl(token):
    target=OUT/'dnl';target.mkdir(parents=True,exist_ok=True)
    base=f'https://api.github.com/repos/{DNL}'
    repo=get(base,token).json(); branch=repo['default_branch'];runs=[]
    # Review the newest artifacts first; show up to 20 retained reports.
    for page in range(1,11):
        batch=get(base+f'/actions/artifacts?per_page=100&page={page}',token).json()['artifacts']
        if not batch:break
        for a in batch:
            if len(runs)>=20:break
            if a['expired'] or not a['name'].startswith('DNL-Report-'):continue
            wr=a.get('workflow_run') or {}
            if wr.get('head_branch')!=branch:continue
            run=get(base+f'/actions/runs/{wr["id"]}',token).json()
            if run['status']!='completed' or run.get('event')=='pull_request':continue
            # Reports may exist even if a later email step failed. Retain that status visibly.
            response=get(base+f'/actions/artifacts/{a["id"]}/zip',token,redirect=False)
            if response.status_code==302:
                location=response.headers['Location']
                if not location.startswith('https://'):raise ValueError('Invalid artifact download address')
                # Signed download uses NO GitHub authorization header.
                response=get(location)
            with zipfile.ZipFile(io.BytesIO(response.content)) as z:
                candidates=[x for x in z.infolist() if Path(x.filename).name.startswith('DNL_') and x.filename.endswith('.xlsx')]
                if not candidates:continue
                chosen=sorted(candidates,key=lambda x:x.filename)[-1]
                if chosen.file_size>50*1024*1024:raise ValueError('DNL Excel exceeds 50 MB limit')
                blob=z.read(chosen); articles=convert_excel(blob)
            day=a['created_at'][:10];name=f'dnl_{day}_{a["id"]}';excel=name+'.xlsx'
            (target/excel).write_bytes(blob)
            save_json(target/(name+'.json'),{'date_short':day,'articles':articles,'excel':excel,'source_run':run['html_url'],'source_conclusion':run['conclusion']})
            runs.append({'date':day,'file':name+'.json','count':len(articles),'excel':excel})
        if len(runs)>=20:break
    save_json(target/'index.json',{'runs':runs})
    print(f'Read {len(runs)} retained DNL Excel reports')

def main():
    token=os.getenv('INTELHUB_READ_TOKEN','').strip()
    if not token:raise RuntimeError('Add INTELHUB_READ_TOKEN to the NEW IntelHub repository secrets. Do not change the source repositories.')
    OUT.mkdir(exist_ok=True)
    for f in ['index.html','integration.js']:shutil.copyfile(ROOT/f,OUT/f)
    sync_radar();sync_dnl(token)
    save_json(OUT/'synced.json',{'synced_at':datetime.now(timezone.utc).isoformat()})
    print('Viewer built. Source repositories were read only; no scans or emails were triggered.')
if __name__=='__main__':main()
